from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# Paleta profesional
COLOR_PERSONAJE_BG = "1B3A4B"
COLOR_PERSONAJE_FONT = "FFFFFF"
COLOR_DIA_BG = "C9A84C"
COLOR_DIA_FONT = "FFFFFF"
COLOR_HEADERS_BG = "4A6FA5"
COLOR_HEADERS_FONT = "FFFFFF"
COLOR_ROW_IMPAR = "FFFFFF"
COLOR_ROW_PAR = "EEF2F7"
COLOR_KBIO_NUEVO = "FFF3CD"
COLOR_DATOS_FONT = "2C3E50"
COLOR_BORDE = "BDC3C7"

FILL_PERSONAJE = PatternFill("solid", fgColor=COLOR_PERSONAJE_BG)
FILL_DIA = PatternFill("solid", fgColor=COLOR_DIA_BG)
FILL_HEADERS = PatternFill("solid", fgColor=COLOR_HEADERS_BG)
FILL_ROW_IMPAR = PatternFill("solid", fgColor=COLOR_ROW_IMPAR)
FILL_ROW_PAR = PatternFill("solid", fgColor=COLOR_ROW_PAR)
FILL_KBIO_NUEVO = PatternFill("solid", fgColor=COLOR_KBIO_NUEVO)

FONT_PERSONAJE = Font(name='Calibri', size=24, bold=True, color=COLOR_PERSONAJE_FONT)
FONT_DIA = Font(name='Calibri', size=14, bold=True, color=COLOR_DIA_FONT)
FONT_HEADER = Font(name='Calibri', size=9, bold=True, color=COLOR_HEADERS_FONT)
FONT_DATOS = Font(name='Calibri', size=10, bold=False, color=COLOR_DATOS_FONT)
FONT_KBIO_VAL = Font(name='Calibri', size=14, bold=True, color=COLOR_PERSONAJE_BG)
FONT_CAP = Font(name='Calibri', size=10, bold=True, color=COLOR_DATOS_FONT)
FONT_X = Font(name='Calibri', size=8, bold=False, color=COLOR_DATOS_FONT)

ALIGN_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
ALIGN_LEFT_WRAP = Alignment(horizontal='left', vertical='top', wrap_text=True)
ALIGN_CENTER_WRAP = Alignment(horizontal='center', wrap_text=True)

THIN_BORDER_SIDE = Side(style='thin', color=COLOR_BORDE)
THIN_BORDER = Border(
    left=THIN_BORDER_SIDE,
    right=THIN_BORDER_SIDE,
    top=THIN_BORDER_SIDE,
    bottom=THIN_BORDER_SIDE,
)
MEDIUM_BORDER_BOTTOM = Border(
    left=THIN_BORDER_SIDE,
    right=THIN_BORDER_SIDE,
    top=THIN_BORDER_SIDE,
    bottom=Side(style='medium', color=COLOR_HEADERS_BG),
)

HEIGHT_PERSONAJE = 32
HEIGHT_DIA = 22
HEIGHT_HEADERS = 15

COL_WIDTHS = {
    'A': 5.88, 'B': 13.0, 'C': 6.44, 'D': 39.44,
    'E': 5.55, 'F': 4.66, 'G': 5.66, 'H': 5.44,
    'I': 70.66, 'J': 40.88,
}

HEADERS = ['KBIO', 'CAP ', 'ESC ', 'LOCACION', 'EXT', 'INT', 'DIA', 'NOC', 'OBSERVACION', 'CONTINUIDAD']


def _set_col_widths(ws):
    for col, width in COL_WIDTHS.items():
        ws.column_dimensions[col].width = width


def _style_cell(cell, font=None, fill=None, alignment=None, border=None):
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if alignment:
        cell.alignment = alignment
    if border:
        cell.border = border


def _write_personaje_header(ws, row, nombre):
    ws.merge_cells(f'A{row}:J{row}')
    cell = ws[f'A{row}']
    cell.value = nombre
    _style_cell(cell, FONT_PERSONAJE, FILL_PERSONAJE, ALIGN_CENTER)
    ws.row_dimensions[row].height = HEIGHT_PERSONAJE
    return row + 1


def _write_dia_header(ws, row, dia_nombre):
    ws.merge_cells(f'A{row}:J{row}')
    cell = ws[f'A{row}']
    cell.value = dia_nombre
    _style_cell(cell, FONT_DIA, FILL_DIA, ALIGN_CENTER_WRAP)
    ws.row_dimensions[row].height = HEIGHT_DIA
    return row + 1


def _write_col_headers(ws, row):
    for col_idx, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=row, column=col_idx)
        cell.value = header
        _style_cell(cell, FONT_HEADER, FILL_HEADERS, ALIGN_CENTER_WRAP, MEDIUM_BORDER_BOTTOM)
    ws.row_dimensions[row].height = HEIGHT_HEADERS
    return row + 1


def _write_data_rows(ws, row_start, filas, prev_kbio=0):
    if not filas:
        return row_start, prev_kbio

    row_end = row_start + len(filas) - 1

    if len(filas) > 1:
        ws.merge_cells(f'A{row_start}:A{row_end}')
        ws.merge_cells(f'J{row_start}:J{row_end}')

    kbio_val = filas[0].kbio if filas[0].kbio is not None else ""
    kbio_cell = ws[f'A{row_start}']
    kbio_cell.value = kbio_val
    kbio_fill = FILL_KBIO_NUEVO if kbio_val and kbio_val != prev_kbio else None
    _style_cell(kbio_cell, FONT_KBIO_VAL, kbio_fill, ALIGN_CENTER, THIN_BORDER)

    cont_cell = ws[f'J{row_start}']
    cont_cell.value = filas[0].continuidad if filas[0].continuidad else ""
    _style_cell(cont_cell, FONT_DATOS, None, ALIGN_CENTER_WRAP, THIN_BORDER)

    last_kbio = kbio_val if kbio_val else prev_kbio

    for i, fila in enumerate(filas):
        r = row_start + i
        row_fill = FILL_ROW_PAR if i % 2 else FILL_ROW_IMPAR
        ws.row_dimensions[r].height = 15.75

        cells_data = [
            (2, fila.capitulo, FONT_CAP, ALIGN_CENTER_WRAP),
            (3, str(fila.escena), FONT_DATOS, ALIGN_CENTER_WRAP),
            (4, fila.locacion, FONT_DATOS, ALIGN_LEFT_WRAP),
            (5, "X" if fila.ext else "", FONT_X, ALIGN_CENTER_WRAP),
            (6, "X" if fila.int_ else "", FONT_X, ALIGN_CENTER_WRAP),
            (7, "X" if fila.dia else "", FONT_X, ALIGN_CENTER_WRAP),
            (8, "X" if fila.noche else "", FONT_X, ALIGN_CENTER_WRAP),
        ]
        for col, val, font, align in cells_data:
            c = ws.cell(row=r, column=col, value=val)
            _style_cell(c, font, row_fill, align, THIN_BORDER)

        obs_text = str(fila.observacion) if fila.observacion else ""
        c = ws.cell(row=r, column=9, value=obs_text)
        _style_cell(c, FONT_DATOS, row_fill, ALIGN_LEFT_WRAP, THIN_BORDER)

        chars = len(obs_text)
        if chars > 280:
            ws.row_dimensions[r].height = 120
        elif chars > 180:
            ws.row_dimensions[r].height = 80
        elif chars > 120:
            ws.row_dimensions[r].height = 57.6
        elif chars > 60:
            ws.row_dimensions[r].height = 28.8
        else:
            ws.row_dimensions[r].height = 15.75

        if len(filas) == 1:
            a_cell = ws.cell(row=r, column=1)
            _style_cell(a_cell, FONT_KBIO_VAL, kbio_fill, ALIGN_CENTER, THIN_BORDER)
            j_cell = ws.cell(row=r, column=10)
            _style_cell(j_cell, FONT_DATOS, row_fill, ALIGN_CENTER_WRAP, THIN_BORDER)

    return row_end + 1, last_kbio


def write_personaje_sheet(ws, nombre, dias_data):
    _set_col_widths(ws)
    row = 1
    prev_kbio_global = 0

    row = _write_personaje_header(ws, row, nombre)

    for dia_nombre, filas in dias_data.items():
        if not filas:
            continue
        row = _write_dia_header(ws, row, dia_nombre)
        row = _write_col_headers(ws, row)
        row, prev_kbio_global = _write_data_rows(ws, row, filas, prev_kbio_global)
        ws.row_dimensions[row].height = 14.4
        row += 1


def write_extras_sheet(ws, dias_data):
    _set_col_widths(ws)
    row = 1
    ws.merge_cells(f'A{row}:J{row}')
    cell = ws[f'A{row}']
    cell.value = "EXTRAS Y FIGURANTES"
    _style_cell(cell, FONT_PERSONAJE, FILL_PERSONAJE, ALIGN_CENTER)
    ws.row_dimensions[row].height = HEIGHT_PERSONAJE
    row += 1

    for dia_nombre, filas in dias_data.items():
        if not filas:
            continue
        row = _write_dia_header(ws, row, dia_nombre)
        row = _write_col_headers(ws, row)
        for i, fila in enumerate(filas):
            row_fill = FILL_ROW_PAR if i % 2 else FILL_ROW_IMPAR
            ws.row_dimensions[row].height = 15.75
            for col, val, font, align in [
                (1, "", FONT_DATOS, ALIGN_CENTER_WRAP),
                (2, fila.capitulo, FONT_CAP, ALIGN_CENTER_WRAP),
                (3, str(fila.escena), FONT_DATOS, ALIGN_CENTER_WRAP),
                (4, fila.locacion, FONT_DATOS, ALIGN_LEFT_WRAP),
                (5, "X" if fila.ext else "", FONT_X, ALIGN_CENTER_WRAP),
                (6, "X" if fila.int_ else "", FONT_X, ALIGN_CENTER_WRAP),
                (7, "X" if fila.dia else "", FONT_X, ALIGN_CENTER_WRAP),
                (8, "X" if fila.noche else "", FONT_X, ALIGN_CENTER_WRAP),
            ]:
                c = ws.cell(row=row, column=col, value=val)
                _style_cell(c, font, row_fill, align, THIN_BORDER)
            obs = str(fila.observacion) if fila.observacion else ""
            c = ws.cell(row=row, column=9, value=obs)
            _style_cell(c, FONT_DATOS, row_fill, ALIGN_LEFT_WRAP, THIN_BORDER)
            c = ws.cell(row=row, column=10, value="")
            _style_cell(c, FONT_DATOS, row_fill, ALIGN_CENTER_WRAP, THIN_BORDER)
            row += 1
        ws.row_dimensions[row].height = 14.4
        row += 1


def write_cambios_sheet(ws, cambios_data):
    row = 1
    for cambio in cambios_data:
        cell = ws.cell(row=row, column=1, value=cambio['personaje'])
        _style_cell(cell, FONT_PERSONAJE, FILL_PERSONAJE, ALIGN_LEFT_WRAP)
        ws.row_dimensions[row].height = HEIGHT_PERSONAJE
        row += 1

        cell = ws.cell(row=row, column=1, value=cambio['resumen'])
        _style_cell(cell, FONT_DIA, FILL_DIA, ALIGN_LEFT_WRAP)
        ws.row_dimensions[row].height = HEIGHT_DIA
        row += 1

        for i, item in enumerate(cambio['items']):
            cell = ws.cell(row=row, column=1, value=item)
            fill = FILL_ROW_PAR if i % 2 else FILL_ROW_IMPAR
            _style_cell(cell, FONT_DATOS, fill, ALIGN_LEFT_WRAP, THIN_BORDER)
            ws.row_dimensions[row].height = 15.75
            row += 1

        row += 2


def _clear_frozen_view(ws):
    """Asegura que la hoja no tenga paneles congelados al desplazarse."""
    if ws.sheet_view.pane is not None:
        ws.sheet_view.pane = None


def write_excel(breakdown, output_path):
    wb = Workbook()
    wb.remove(wb.active)

    for nombre, dias in breakdown['personajes'].items():
        safe_name = nombre[:31].replace('/', '-').replace('\\', '-').replace('?', '').replace('*', '').replace('[', '').replace(']', '').replace(':', '')
        ws = wb.create_sheet(title=safe_name)
        write_personaje_sheet(ws, nombre, dias)

    if breakdown.get('extras'):
        ws = wb.create_sheet(title="EXTRAS Y FIGURANTES")
        write_extras_sheet(ws, breakdown['extras'])

    if breakdown.get('cambios'):
        ws = wb.create_sheet(title="CAMBIOS")
        write_cambios_sheet(ws, breakdown['cambios'])

    for ws in wb.worksheets:
        _clear_frozen_view(ws)

    wb.save(output_path)
    return output_path
